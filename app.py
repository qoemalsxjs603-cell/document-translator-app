
import io
import os
import re
import json
import hashlib
import tempfile
from pathlib import Path

import pandas as pd
import streamlit as st
from dotenv import load_dotenv
from google import genai

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from docx import Document
from openpyxl import load_workbook

load_dotenv()

SUPPORTED_EXTS = [".pptx", ".docx", ".xlsx"]

LANGUAGE_OPTIONS = {
    "중국어": "Simplified Chinese",
    "영어": "English",
}

DEFAULT_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")


def get_secret_or_env(key, default=None):
    try:
        if key in st.secrets:
            return st.secrets[key]
    except Exception:
        pass
    return os.getenv(key, default)


def has_translatable_text(text: str) -> bool:
    if text is None:
        return False
    s = str(text).strip()
    if not s:
        return False
    if s.startswith("="):
        return False
    return bool(re.search(r"[A-Za-z가-힣一-龥ぁ-ゔァ-ヴー]", s))


def normalize_text(text: str) -> str:
    return str(text).replace("\r\n", "\n").replace("\r", "\n")


def make_hash(text: str, target_language: str, glossary_text: str, style: str) -> str:
    raw = f"{target_language}|{style}|{glossary_text}|{text}"
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def safe_filename(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "_", name)


def make_output_name(original_name: str, target_label: str) -> str:
    p = Path(original_name)
    lang = "CN" if target_label == "중국어" else "EN"
    return safe_filename(f"{p.stem}_translated_{lang}{p.suffix}")


def extract_json_array(text: str):
    text = text.strip()
    try:
        return json.loads(text)
    except Exception:
        pass

    start = text.find("[")
    end = text.rfind("]")
    if start >= 0 and end > start:
        return json.loads(text[start:end+1])

    raise ValueError("모델 응답에서 JSON 배열을 찾지 못했습니다.")


@st.cache_resource
def get_client():
    api_key = get_secret_or_env("GEMINI_API_KEY")
    if not api_key:
        raise ValueError("GEMINI_API_KEY가 설정되어 있지 않습니다. Streamlit Secrets에 API 키를 입력하세요.")
    return genai.Client(api_key=api_key)


def build_glossary_text(glossary_file):
    if glossary_file is None:
        return ""

    name = glossary_file.name.lower()
    try:
        if name.endswith(".xlsx"):
            df = pd.read_excel(glossary_file)
        elif name.endswith(".csv"):
            df = pd.read_csv(glossary_file)
        else:
            return ""
    except Exception:
        return ""

    if df.empty:
        return ""

    cols = list(df.columns)
    if len(cols) < 2:
        return ""

    pairs = []
    for _, row in df.iterrows():
        src = str(row[cols[0]]).strip()
        dst = str(row[cols[1]]).strip()
        if src and dst and src.lower() != "nan" and dst.lower() != "nan":
            pairs.append(f"{src} => {dst}")

    return "\n".join(pairs[:1000])


def translate_batch(texts, target_language, model, style, glossary_text):
    if not texts:
        return []

    client = get_client()

    glossary_rule = ""
    if glossary_text.strip():
        glossary_rule = f"""
Glossary:
Use the following glossary strictly when applicable.
{glossary_text}
""".strip()

    system_prompt = f"""
You are a professional business document translator.

Translate each source string into {target_language}.

Rules:
- Return ONLY a valid JSON array of strings.
- Output array length must equal input array length.
- Do not add explanations.
- Preserve line breaks as much as possible.
- Preserve numbers, dates, names, brand names, URLs, email addresses, and placeholders.
- Preserve bullet symbols when possible.
- Keep titles concise.
- If a phrase is already in the target language, keep it natural.
- If the source contains mixed Korean/English, translate Korean and keep necessary English terms.

Style:
{style}

{glossary_rule}
""".strip()

    prompt = f"""
{system_prompt}

Input JSON array:
{json.dumps(texts, ensure_ascii=False)}
""".strip()

    response = client.models.generate_content(
        model=model,
        contents=prompt,
    )

    arr = extract_json_array(response.text)

    if len(arr) != len(texts):
        raise ValueError(f"번역 결과 개수 불일치: 원문 {len(texts)}개 / 번역 {len(arr)}개")

    return ["" if x is None else str(x) for x in arr]


def translate_texts_with_cache(texts, target_language, model, style, glossary_text, batch_size, progress=None):
    if "translation_cache" not in st.session_state:
        st.session_state["translation_cache"] = {}

    cache = st.session_state["translation_cache"]

    unique = []
    seen = set()

    for t in texts:
        if not has_translatable_text(t):
            continue
        t = normalize_text(t)
        key = make_hash(t, target_language, glossary_text, style)
        if key in seen:
            continue
        seen.add(key)
        if key not in cache:
            unique.append(t)

    total = len(unique)
    done = 0

    for i in range(0, total, batch_size):
        batch = unique[i:i+batch_size]
        translated = translate_batch(batch, target_language, model, style, glossary_text)
        for src, dst in zip(batch, translated):
            key = make_hash(src, target_language, glossary_text, style)
            cache[key] = dst

        done += len(batch)
        if progress is not None:
            progress.progress(min(done / max(total, 1), 1.0), text=f"번역 중... {done}/{total}")

    result_map = {}
    for t in texts:
        if has_translatable_text(t):
            t_norm = normalize_text(t)
            key = make_hash(t_norm, target_language, glossary_text, style)
            if key in cache:
                result_map[t_norm] = cache[key]

    return result_map


def make_log_excel(rows, output_path):
    df = pd.DataFrame(rows)
    df.to_excel(output_path, index=False)


def iter_ppt_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub in iter_ppt_shapes(shape.shapes):
                yield sub


def collect_pptx_texts(path):
    prs = Presentation(path)
    texts = []
    meta = []

    for s_idx, slide in enumerate(prs.slides, start=1):
        for shape in iter_ppt_shapes(slide.shapes):
            try:
                if getattr(shape, "has_table", False):
                    for r_idx, row in enumerate(shape.table.rows, start=1):
                        for c_idx, cell in enumerate(row.cells, start=1):
                            txt = normalize_text(cell.text)
                            if has_translatable_text(txt):
                                texts.append(txt)
                                meta.append({"file_type": "pptx", "location": f"slide {s_idx} table r{r_idx}c{c_idx}", "source": txt})
                elif getattr(shape, "has_text_frame", False):
                    txt = normalize_text(shape.text)
                    if has_translatable_text(txt):
                        texts.append(txt)
                        meta.append({"file_type": "pptx", "location": f"slide {s_idx} text_box", "source": txt})
            except Exception:
                pass

    return texts, meta


def apply_to_text_frame(text_frame, translated):
    if text_frame.paragraphs and text_frame.paragraphs[0].runs:
        text_frame.paragraphs[0].runs[0].text = translated
        for p in text_frame.paragraphs:
            for r in p.runs[1:]:
                r.text = ""
        for p in text_frame.paragraphs[1:]:
            for r in p.runs:
                r.text = ""
    else:
        text_frame.text = translated


def translate_pptx(input_path, translation_map, output_path):
    prs = Presentation(input_path)

    for slide in prs.slides:
        for shape in iter_ppt_shapes(slide.shapes):
            try:
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        for cell in row.cells:
                            src = normalize_text(cell.text)
                            if src in translation_map:
                                apply_to_text_frame(cell.text_frame, translation_map[src])
                elif getattr(shape, "has_text_frame", False):
                    src = normalize_text(shape.text)
                    if src in translation_map:
                        apply_to_text_frame(shape.text_frame, translation_map[src])
            except Exception:
                pass

    prs.save(output_path)


def collect_docx_texts(path):
    doc = Document(path)
    texts = []
    meta = []

    def collect_paragraphs(paragraphs, loc):
        for i, p in enumerate(paragraphs, start=1):
            txt = normalize_text(p.text)
            if has_translatable_text(txt):
                texts.append(txt)
                meta.append({"file_type": "docx", "location": f"{loc} paragraph {i}", "source": txt})

    def collect_tables(tables, loc):
        for t_idx, table in enumerate(tables, start=1):
            for r_idx, row in enumerate(table.rows, start=1):
                for c_idx, cell in enumerate(row.cells, start=1):
                    collect_paragraphs(cell.paragraphs, f"{loc} table {t_idx} r{r_idx}c{c_idx}")
                    collect_tables(cell.tables, f"{loc} nested table {t_idx}")

    collect_paragraphs(doc.paragraphs, "body")
    collect_tables(doc.tables, "body")

    for s_idx, section in enumerate(doc.sections, start=1):
        collect_paragraphs(section.header.paragraphs, f"section {s_idx} header")
        collect_tables(section.header.tables, f"section {s_idx} header")
        collect_paragraphs(section.footer.paragraphs, f"section {s_idx} footer")
        collect_tables(section.footer.tables, f"section {s_idx} footer")

    return texts, meta


def apply_to_paragraph(paragraph, translated):
    if paragraph.runs:
        paragraph.runs[0].text = translated
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(translated)


def translate_docx(input_path, translation_map, output_path):
    doc = Document(input_path)

    def apply_paragraphs(paragraphs):
        for p in paragraphs:
            src = normalize_text(p.text)
            if src in translation_map:
                apply_to_paragraph(p, translation_map[src])

    def apply_tables(tables):
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    apply_paragraphs(cell.paragraphs)
                    apply_tables(cell.tables)

    apply_paragraphs(doc.paragraphs)
    apply_tables(doc.tables)

    for section in doc.sections:
        apply_paragraphs(section.header.paragraphs)
        apply_tables(section.header.tables)
        apply_paragraphs(section.footer.paragraphs)
        apply_tables(section.footer.tables)

    doc.save(output_path)


def collect_xlsx_texts(path):
    wb = load_workbook(path)
    texts = []
    meta = []

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if isinstance(val, str):
                    txt = normalize_text(val)
                    if has_translatable_text(txt):
                        texts.append(txt)
                        meta.append({"file_type": "xlsx", "location": f"{ws.title}!{cell.coordinate}", "source": txt})

    return texts, meta


def translate_xlsx(input_path, translation_map, output_path):
    wb = load_workbook(input_path)

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if isinstance(val, str):
                    src = normalize_text(val)
                    if src in translation_map:
                        cell.value = translation_map[src]

    wb.save(output_path)


def main():
    st.set_page_config(page_title="문서 번역기 Gemini", page_icon="🌐", layout="wide")

    st.title("문서 번역기 Gemini")
    st.caption("PPTX / DOCX / XLSX 파일을 업로드하면 중국어 또는 영어로 번역한 파일을 생성합니다.")

    with st.sidebar:
        st.header("번역 설정")
        target_label = st.selectbox("번역 언어", ["중국어", "영어"])
        model = st.text_input("Gemini 모델", value=get_secret_or_env("GEMINI_MODEL", DEFAULT_MODEL))

        style_label = st.selectbox(
            "번역 스타일",
            [
                "실무 제출용 자연 번역",
                "원문 구조 유지 직역",
                "외부 발표용 마케팅 문체",
            ],
        )

        style_map = {
            "실무 제출용 자연 번역": "Business document style. Natural, clear, faithful, and suitable for internal work submission.",
            "원문 구조 유지 직역": "Literal translation. Keep the source wording and structure as much as possible while remaining understandable.",
            "외부 발표용 마케팅 문체": "Polished marketing style. Natural and suitable for external presentations.",
        }

        batch_size = st.slider("번역 배치 크기", 5, 60, 20, 5)

        st.divider()
        st.subheader("선택 기능")
        use_glossary = st.checkbox("용어집 사용", value=False)
        glossary_file = None
        if use_glossary:
            glossary_file = st.file_uploader("용어집 업로드 (.xlsx 또는 .csv)", type=["xlsx", "csv"])

        make_log = st.checkbox("번역 로그 엑셀 생성", value=True)

        if st.button("세션 번역 캐시 초기화"):
            st.session_state["translation_cache"] = {}
            st.success("캐시를 초기화했습니다.")

    uploaded_files = st.file_uploader(
        "문서 파일 업로드",
        type=["pptx", "docx", "xlsx"],
        accept_multiple_files=True,
        help="여러 파일을 한 번에 업로드할 수 있습니다.",
    )

    if not get_secret_or_env("GEMINI_API_KEY"):
        st.warning("GEMINI_API_KEY가 설정되어 있지 않습니다. Streamlit Secrets에 API 키를 입력하세요.")

    if uploaded_files:
        st.write("업로드된 파일")
        st.dataframe(
            pd.DataFrame([{"파일명": f.name, "크기(KB)": round(len(f.getvalue()) / 1024, 1)} for f in uploaded_files]),
            use_container_width=True,
        )

    if uploaded_files and st.button("번역 시작", type="primary"):
        target_language = LANGUAGE_OPTIONS[target_label]
        style = style_map[style_label]
        glossary_text = build_glossary_text(glossary_file) if use_glossary else ""

        results = []

        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)

            for file_index, uploaded in enumerate(uploaded_files, start=1):
                ext = Path(uploaded.name).suffix.lower()
                if ext not in SUPPORTED_EXTS:
                    st.error(f"지원하지 않는 파일 형식입니다: {uploaded.name}")
                    continue

                st.subheader(f"{file_index}. {uploaded.name}")

                input_path = tmp / safe_filename(uploaded.name)
                input_path.write_bytes(uploaded.getvalue())

                output_name = make_output_name(uploaded.name, target_label)
                output_path = tmp / output_name
                log_path = tmp / f"{Path(output_name).stem}_translation_log.xlsx"

                try:
                    with st.spinner("텍스트 추출 중..."):
                        if ext == ".pptx":
                            texts, meta = collect_pptx_texts(input_path)
                        elif ext == ".docx":
                            texts, meta = collect_docx_texts(input_path)
                        elif ext == ".xlsx":
                            texts, meta = collect_xlsx_texts(input_path)

                    unique_count = len(set(texts))
                    st.write(f"추출 텍스트: {len(texts)}개 / 중복 제거 후: {unique_count}개")

                    progress = st.progress(0, text="번역 준비 중...")
                    translation_map = translate_texts_with_cache(
                        texts=texts,
                        target_language=target_language,
                        model=model,
                        style=style,
                        glossary_text=glossary_text,
                        batch_size=batch_size,
                        progress=progress,
                    )

                    with st.spinner("문서에 번역 삽입 중..."):
                        if ext == ".pptx":
                            translate_pptx(input_path, translation_map, output_path)
                        elif ext == ".docx":
                            translate_docx(input_path, translation_map, output_path)
                        elif ext == ".xlsx":
                            translate_xlsx(input_path, translation_map, output_path)

                    log_rows = []
                    for item in meta:
                        src = normalize_text(item["source"])
                        log_rows.append({
                            "file": uploaded.name,
                            "file_type": item["file_type"],
                            "location": item["location"],
                            "source": src,
                            "translation": translation_map.get(src, ""),
                        })

                    if make_log:
                        make_log_excel(log_rows, log_path)

                    results.append({
                        "name": output_name,
                        "bytes": output_path.read_bytes(),
                        "log_name": log_path.name if make_log else None,
                        "log_bytes": log_path.read_bytes() if make_log else None,
                    })

                    st.success("완료")

                    st.download_button(
                        "번역 파일 다운로드",
                        data=output_path.read_bytes(),
                        file_name=output_name,
                        mime="application/octet-stream",
                        key=f"download_{file_index}",
                    )

                    if make_log:
                        st.download_button(
                            "번역 로그 엑셀 다운로드",
                            data=log_path.read_bytes(),
                            file_name=log_path.name,
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            key=f"log_{file_index}",
                        )

                    with st.expander("미리보기: 번역 로그 상위 30개"):
                        st.dataframe(pd.DataFrame(log_rows).head(30), use_container_width=True)

                except Exception as e:
                    st.error(f"오류 발생: {uploaded.name}")
                    st.exception(e)

            if len(results) > 1:
                zip_buffer = io.BytesIO()
                import zipfile
                with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as z:
                    for r in results:
                        z.writestr(r["name"], r["bytes"])
                        if r["log_name"] and r["log_bytes"]:
                            z.writestr(r["log_name"], r["log_bytes"])
                zip_buffer.seek(0)

                st.divider()
                st.download_button(
                    "전체 결과 ZIP 다운로드",
                    data=zip_buffer.getvalue(),
                    file_name=f"translated_documents_{'CN' if target_label == '중국어' else 'EN'}.zip",
                    mime="application/zip",
                )

    with st.expander("지원 범위와 한계"):
        st.markdown(
            """
            **지원**
            - PPTX: 텍스트 박스, 표, 일부 그룹 오브젝트
            - DOCX: 본문, 표, 헤더, 푸터
            - XLSX: 일반 텍스트 셀
            - 다중 파일 업로드
            - 번역 로그 엑셀 저장
            - 세션 내 번역 캐시
            - 선택형 용어집 업로드

            **한계**
            - 이미지로 박힌 글자는 번역하지 못합니다.
            - 일부 SmartArt / WordArt는 누락될 수 있습니다.
            - PPT는 번역 후 글자 넘침, 줄바꿈, 폰트 검수가 필요할 수 있습니다.
            - 암호화된 파일은 지원하지 않습니다.
            """
        )


if __name__ == "__main__":
    main()
